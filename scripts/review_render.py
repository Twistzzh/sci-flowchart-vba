# -*- coding: utf-8 -*-
"""第二轮评审的程序化几何检查（对 build_ppt.py 产出的 .pptx）。

用法:
    python review_render.py <输出目录|flowchart.pptx路径>

检查项（对应 references/review_rubric.md 评审二）:
  1. 文本溢出   PIL 按真实字体量宽（无字体文件时退化为系数估算），
                文字溢出所属节点框即 FAIL；单行超宽但放得下高度 ->
                依赖自动折行，WARN。
  2. 悬空连线   连线端点必须落在某节点边界带内，否则 FAIL。
  3. 连线穿模   连线线段穿入非容器节点的严格内部 -> WARN（贴边不算）。
  4. 节点重叠   两个**都有文字**的实体节点明显交叠且非包含 -> WARN
                （装饰衬框/底色块无文字，重叠属设计行为，跳过）。
  5. 出界       形状超出幻灯片边界 -> FAIL。
  6. 字号观感   所有节点"文本宽/可用宽"中位数过低 -> 提示字体可能过小。

退出码: 0 = PASS（允许 WARN），1 = FAIL。文本型结论打印到 stdout。
"""
import math
import os
import sys

FALLBACK_K = 0.55   # 无字体文件时的估算系数（em/字符）
K_LINE = 1.25       # 行高系数
TOL_EDGE = 0.06     # 端点贴边容差（英寸）
SHRINK_IN = 0.02    # "严格内部"判定收缩量（英寸）
TOL_OVERLAP = 0.15  # 节点重叠: 交叠面积 / 较小面积
CONTAINER_MIN = 3   # bbox 内包含 >=3 个其他节点中心 -> 视为容器面板

FONT_FILES = {
    "times new roman": "times.ttf",
    "times": "times.ttf",
    "arial": "arial.ttf",
    "calibri": "calibri.ttf",
    "cambria": "cambria.ttf",
    "georgia": "georgia.ttf",
    "segoe ui": "segoeui.ttf",
    "helvetica": "arial.ttf",
    "microsoft yahei": "msyh.ttc",
    "yahei": "msyh.ttc",
    "msyh": "msyh.ttc",
    "simsun": "simsun.ttc",
    "simhei": "simhei.ttf",
}
_font_cache = {}
_pil_ok = None


def _load_font(font_name, pt, bold=False):
    global _pil_ok
    if _pil_ok is False:
        return None
    try:
        from PIL import ImageFont
    except ImportError:
        _pil_ok = False
        return None
    key = (font_name.lower(), int(pt * 4), bold)
    if key in _font_cache:
        return _font_cache[key]
    fn = FONT_FILES.get(font_name.lower())
    if not fn:
        _pil_ok = False
        return None
    # bold CJK: prefer the bold face when available (msyh.ttc -> msyhbd.ttc)
    if bold and fn == "msyh.ttc":
        bpath = os.path.join(os.environ.get("WINDIR", r"C:\Windows"),
                             "Fonts", "msyhbd.ttc")
        if os.path.isfile(bpath):
            fn = "msyhbd.ttc"
    path = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", fn)
    if not os.path.isfile(path):
        _pil_ok = False
        return None
    try:
        f = ImageFont.truetype(path, size=max(8, int(pt * 4)))
    except Exception:
        _pil_ok = False
        return None
    _font_cache[key] = f
    return f


def _cjk_subst(font_name):
    """If font_name is a Latin-only family but the text is CJK, return the
    matching CJK face name so width measurement uses a real CJK font."""
    f = (font_name or "").lower()
    if any(k in f for k in ("yahei", "msyh", "arial", "helvetica", "segoe", "sans")):
        return "microsoft yahei"
    if any(k in f for k in ("simsun", "宋", "sun", "serif", "times")):
        return "simsun"
    if any(k in f for k in ("simhei", "黑", "hei")):
        return "simhei"
    return None


def is_cjk(ch):
    return ('\u3000' <= ch <= '\u303f' or '\u3400' <= ch <= '\u9fff'
            or '\uff00' <= ch <= '\uffef')


def text_width_in(text, font_name, pt, bold=False):
    """文本渲染宽度（英寸）。PIL 真实量宽优先，缺失时退化估算（CJK 感知）。"""
    f = _load_font(font_name, pt, bold)
    if f is None and any(is_cjk(c) for c in text):
        # Latin font won't measure CJK; substitute the matching CJK face
        sub = _cjk_subst(font_name)
        if sub:
            f = _load_font(sub, pt, bold)
    if f is not None:
        try:
            return f.getlength(text) / 4.0 / 72.0
        except Exception:
            pass
    # 无字体文件时的估算：CJK 字形约 1.0em，西文约 0.55em，空格约 0.3em
    w = 0.0
    for ch in text:
        if ch == ' ':
            w += 0.3 * pt
        elif is_cjk(ch):
            w += 1.0 * pt
        else:
            w += FALLBACK_K * pt
    return w / 72.0


def emu2in(v):
    return v / 914400.0


def bbox(shp):
    return (emu2in(shp.left), emu2in(shp.top),
            emu2in(shp.left + shp.width), emu2in(shp.top + shp.height))


def inside(px, py, b, tol=TOL_EDGE):
    x0, y0, x1, y1 = b
    return (x0 - tol <= px <= x1 + tol) and (y0 - tol <= py <= y1 + tol)


def seg_hits_inner(p0, p1, b, shrink=SHRINK_IN):
    """线段是否穿入 bbox 收缩 shrink 后的严格内部（贴边不算穿越）。"""
    x0, y0, x1, y1 = b[0] + shrink, b[1] + shrink, b[2] - shrink, b[3] - shrink
    if x0 >= x1 or y0 >= y1:
        return False
    (ax, ay), (bx, by) = p0, p1
    if max(ax, bx) < x0 or min(ax, bx) > x1:
        return False
    if max(ay, by) < y0 or min(ay, by) > y1:
        return False
    for t in (i / 24.0 for i in range(25)):
        px, py = ax + (bx - ax) * t, ay + (by - ay) * t
        if x0 < px < x1 and y0 < py < y1:
            return True
    return False


def text_metrics(tf, default_pt=18.0):
    """返回 (lines, pt, font, known) —— 行列表、字号、字体名、字号是否可读。"""
    lines, sizes, fonts = [], [], []
    for para in tf.paragraphs:
        t = "".join(r.text for r in para.runs)
        if not t and para.runs:
            t = " "
        if t:
            lines.append(t)
        for r in para.runs:
            if r.font.size is not None:
                sizes.append(r.font.size.pt)
            if r.font.name:
                fonts.append(r.font.name)
    pt = max(sizes) if sizes else None
    font = fonts[0] if fonts else "Arial"
    return lines, (pt or default_pt), font, pt is not None


def review(pptx_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(pptx_path)
    errs, warns, infos = [], [], []

    slide = prs.slides[0]
    SW, SH = emu2in(prs.slide_width), emu2in(prs.slide_height)

    nodes, conns = [], []
    for shp in slide.shapes:
        is_line = (shp.shape_type == MSO_SHAPE_TYPE.LINE)
        if is_line:
            try:
                p0 = (emu2in(shp.begin_x), emu2in(shp.begin_y))
                p1 = (emu2in(shp.end_x), emu2in(shp.end_y))
                conns.append((p0, p1))
            except Exception:
                b = bbox(shp)
                conns.append(((b[0], b[1]), (b[2], b[3])))
            continue
        b = bbox(shp)
        tf = shp.text_frame if shp.has_text_frame else None
        label = tf.text.replace("\n", " / ")[:36] if tf is not None else ""
        nodes.append({"id": shp.shape_id,
                      "name": "%s [%s]" % (shp.shape_id, label),
                      "b": b, "tf": tf, "text": bool(label.strip())})

    def is_container(n):
        cx = (n["b"][0] + n["b"][2]) / 2.0
        cy = (n["b"][1] + n["b"][3]) / 2.0
        cnt = sum(1 for m in nodes if m is not n and
                  inside(cx, cy, m["b"], tol=0.0))
        return cnt >= CONTAINER_MIN

    containers = [n for n in nodes if is_container(n)]
    solids = [n for n in nodes if not is_container(n)]

    # ---- 1) 文本溢出 + 字号观感 -------------------------------------------
    ratios = []
    for n in solids:
        if n["tf"] is None:
            continue
        tf = n["tf"]
        lines, pt, font, known = text_metrics(tf)
        if not lines:
            continue
        if not known:
            warns.append("%s: 字号未显式设置（读不到 run.font.size）" % n["name"])
        ml = emu2in(tf.margin_left) if tf.margin_left is not None else 0.05
        mr = emu2in(tf.margin_right) if tf.margin_right is not None else 0.05
        mt = emu2in(tf.margin_top) if tf.margin_top is not None else 0.02
        mb = emu2in(tf.margin_bottom) if tf.margin_bottom is not None else 0.02
        avail_w = (n["b"][2] - n["b"][0]) - (ml + mr)
        avail_h = (n["b"][3] - n["b"][1]) - (mt + mb)
        if avail_w <= 0 or avail_h <= 0:
            continue
        line_ws = [text_width_in(l, font, pt) for l in lines]
        # 单行字形实际高度约等于 pt（vertical-anchor middle 下安全）；
        # 多行才按行距 1.25 累加
        h_coeff = 1.0 if len(lines) == 1 else K_LINE
        total_h = len(lines) * pt * h_coeff / 72.0
        ratios.append(max(line_ws) / avail_w)
        over = [(l, w) for l, w in zip(lines, line_ws) if w > avail_w]
        if over:
            if len(lines) == 1 and total_h <= avail_h:
                wrap_n = int(math.ceil(line_ws[0] / avail_w))
                if wrap_n * pt * K_LINE / 72.0 > avail_h:
                    errs.append("%s: 字号 %gpt 过大，'%s' 自动折行 %d 行仍溢出节点"
                                % (n["name"], pt, over[0][0][:30], wrap_n))
                else:
                    warns.append("%s: '%s' 超宽依赖自动折行（应用 vbLf 显式分行）"
                                 % (n["name"], over[0][0][:30]))
            else:
                for l, w in over:
                    errs.append("%s: 字号 %gpt 文字 '%s' 溢出节点宽（est %.2f in > 可用 %.2f in）"
                                % (n["name"], pt, l[:30], w, avail_w))
        if total_h > avail_h + 0.02:
            errs.append("%s: 文本总高 %.2f in 超出节点可用高 %.2f in（字号过大或行数过多）"
                        % (n["name"], total_h, avail_h))
    if ratios:
        med = sorted(ratios)[len(ratios) // 2]
        if med < 0.35:
            warns.append("全部节点文字宽/框宽中位数 %.2f < 0.35：字体可能整体偏小（对照原图比例）" % med)
        else:
            infos.append("文字宽/框宽中位数 %.2f（0.35~1.0 为合理区间）" % med)

    # ---- 2) 悬空连线 -------------------------------------------------------
    # 折线拐点：某端点若与另一条连线的端点重合（距离 < SNAP），说明它是
    # 多段折线的中间拐点而非自由端，不判悬空。真悬空 = 既不接节点也不接线。
    SNAP = 0.03
    endpoints = [p for seg in conns for p in seg]

    def at_joint(p):
        return any(q is not p and p is not q and
                   abs(p[0] - q[0]) < SNAP and abs(p[1] - q[1]) < SNAP
                   for q in endpoints)

    for i, (p0, p1) in enumerate(conns):
        for tag, p in (("起点", p0), ("终点", p1)):
            if not any(inside(p[0], p[1], m["b"]) for m in nodes) and \
               not at_joint(p):
                errs.append("连线#%d %s (%.2f,%.2f) 悬空：不在任何节点边界上"
                            % (i, tag, p[0], p[1]))
                break

    # ---- 3) 连线穿模（贴边不算；端点所在节点跳过） -------------------------
    for i, (p0, p1) in enumerate(conns):
        for m in solids:
            if inside(p0[0], p0[1], m["b"], tol=0.0) or \
               inside(p1[0], p1[1], m["b"], tol=0.0):
                continue
            if seg_hits_inner(p0, p1, m["b"]):
                warns.append("连线#%d 穿过节点 %s 内部" % (i, m["name"]))
                break

    # ---- 4) 节点意外重叠（仅双方都有文字时才算事故） -----------------------
    for i in range(len(solids)):
        for j in range(i + 1, len(solids)):
            a, c = solids[i], solids[j]
            if not (a["text"] and c["text"]):
                continue  # 无文字的横幅/衬框/底色块重叠属设计行为
            b2a, b2b = a["b"], c["b"]
            ox = min(b2a[2], b2b[2]) - max(b2a[0], b2b[0])
            oy = min(b2a[3], b2b[3]) - max(b2a[1], b2b[1])
            if ox <= 0.02 or oy <= 0.02:
                continue
            if (b2a[0] <= b2b[0] and b2a[1] <= b2b[1] and b2a[2] >= b2b[2] and b2a[3] >= b2b[3]) or \
               (b2b[0] <= b2a[0] and b2b[1] <= b2a[1] and b2b[2] >= b2a[2] and b2b[3] >= b2a[3]):
                continue  # 包含关系（双框/衬底）合法
            area_min = min((b2a[2] - b2a[0]) * (b2a[3] - b2a[1]),
                           (b2b[2] - b2b[0]) * (b2b[3] - b2b[1]))
            if area_min > 0 and (ox * oy) / area_min >= TOL_OVERLAP:
                warns.append("节点 %s 与 %s 意外重叠（交叠 %.0f%%）"
                             % (a["name"], c["name"], 100 * (ox * oy) / area_min))

    # ---- 5) 出界 ------------------------------------------------------------
    for n in nodes:
        b = n["b"]
        if b[0] < -0.01 or b[1] < -0.01 or b[2] > SW + 0.01 or b[3] > SH + 0.01:
            errs.append("节点 %s 超出幻灯片边界 %s"
                        % (n["name"], (round(SW, 2), round(SH, 2))))
    _ = containers
    return errs, warns, infos, len(nodes), len(conns)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    target = sys.argv[1]
    if os.path.isdir(target):
        # 过滤 Office 锁文件（~$xxx.pptx 不是有效包，且 mtime 往往最新）
        cand = [f for f in os.listdir(target)
                if f.endswith(".pptx") and not f.startswith("~$")]
        if not cand:
            print("FAIL: 目录里没有 .pptx，请先跑 build_ppt.py")
            return 1
        # 多个 pptx 时取修改时间最新的（上轮产物可能被锁定无法覆盖）
        target = os.path.join(target, max(cand, key=lambda f: os.path.getmtime(os.path.join(target, f))))
    if not os.path.isfile(target):
        print("FAIL: 找不到 %s" % target)
        return 1
    print("review_render: %s" % target)
    try:
        errs, warns, infos, nn, nc = review(target)
    except ImportError:
        print("FAIL: 需要 python-pptx，请用隔离 venv 运行:")
        print("  C:/Users/Hello/.workbuddy/binaries/python/envs/default/Scripts/python.exe "
              + sys.argv[0])
        return 2
    print("shapes: %d nodes / %d connector segments" % (nn, nc))
    for line in errs:
        print("  [ERR ] %s" % line)
    for line in warns:
        print("  [WARN] %s" % line)
    for line in infos:
        print("  [INFO] %s" % line)
    if errs:
        print("RESULT: FAIL (%d err / %d warn)" % (len(errs), len(warns)))
        print("按 references/review_rubric.md 评审二处置：修 .bas -> 重跑 build_ppt.py -> 再评审")
        return 1
    print("RESULT: PASS (%d warn)" % len(warns))
    print("程序化检查通过。仍需多模态对照 preview_compare.png 做风格终审（评审一 + 评审二目测项）。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
